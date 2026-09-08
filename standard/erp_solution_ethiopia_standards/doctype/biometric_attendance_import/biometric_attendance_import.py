# Copyright (c) 2026
# Biometric Attendance Import - controller
#
# Place this file at:
#   <your_app>/<your_app>/hr/doctype/biometric_attendance_import/biometric_attendance_import.py

import datetime
import re

import frappe
import openpyxl
from frappe.model.document import Document
from frappe.utils import escape_html, getdate


# ---------------------------------------------------------------------------
# Attendance status logic
# ---------------------------------------------------------------------------

def determine_status(exception_text, regular_hours, overtime_hours, total_hours):
    """
    Determine Attendance status.

    Current default rule:
        Absence -> Absent
        Everything else -> Present

    Replace this function when the exact business rule is confirmed.
    """

    exception_text = (exception_text or "").strip()

    if exception_text.lower() == "absence":
        return "Absent"

    return "Present"


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------

# Expected employee header examples:
#   John Doe (123 : Sales)
#
# The regex captures:
#   1. Employee name
#   2. Device ID
#   3. Department
#
EMP_HEADER_RE = re.compile(
    r"^\s*(.+?)\s*\((\d+)\s*:\s*(.*?)\)\s*$"
)

# Expected date:
#   2026/08/21
DATE_RE = re.compile(
    r"^\s*(\d{4}/\d{2}/\d{2})\s*$"
)


def _to_hours(value):
    """
    Convert an Excel time value into decimal hours.

    Examples:
        08:30 -> 8.5
        04:15 -> 4.25

    Returns 0.0 for unsupported/empty values.
    """

    if isinstance(value, datetime.datetime):
        value = value.time()

    if isinstance(value, datetime.time):
        return round(
            value.hour
            + value.minute / 60
            + value.second / 3600,
            2,
        )

    if isinstance(value, datetime.timedelta):
        return round(value.total_seconds() / 3600, 2)

    if isinstance(value, (int, float)):
        # Excel stores times as a fraction of a day.
        # Example: 0.5 = 12 hours.
        if 0 <= value <= 1:
            return round(value * 24, 2)

        return round(float(value), 2)

    return 0.0


def _parse_workbook(file_path, from_date, to_date):
    """
    Read the biometric Excel file.

    Returns one dictionary per employee/day.

    Duplicate raw-punch rows are skipped when column F is populated on
    what should otherwise be the aggregated date row.
    """

    if not from_date or not to_date:
        frappe.throw("Attendance From Date and Attendance To Date are required.")

    from_date = getdate(from_date)
    to_date = getdate(to_date)

    if from_date > to_date:
        frappe.throw(
            "Attendance From Date cannot be later than Attendance To Date."
        )

    wb = openpyxl.load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    ws = wb.active

    rows_out = []

    current_name = None
    current_device_id = None
    current_department = None

    for row_number in range(1, ws.max_row + 1):

        # Column A
        a = ws.cell(row_number, 1).value

        if a is None:
            continue

        a_str = str(a).strip()

        # ---------------------------------------------------------------
        # Employee header
        # ---------------------------------------------------------------

        employee_match = EMP_HEADER_RE.match(a_str)

        if employee_match:
            current_name = employee_match.group(1).strip()
            current_device_id = employee_match.group(2).strip()
            current_department = employee_match.group(3).strip()

            continue

        # ---------------------------------------------------------------
        # Date row
        # ---------------------------------------------------------------

        date_match = DATE_RE.match(a_str)

        if not date_match:
            continue

        # If there is no employee header yet, the date cannot be matched.
        if not current_device_id:
            continue

        # ---------------------------------------------------------------
        # Skip duplicate raw punch detail rows.
        #
        # Column F = Name in the biometric export.
        # A normal aggregated date row should have this empty.
        # ---------------------------------------------------------------

        if ws.cell(row_number, 6).value is not None:
            continue

        date_str = date_match.group(1)

        try:
            attendance_date = datetime.datetime.strptime(
                date_str,
                "%Y/%m/%d",
            ).date()
        except ValueError:
            continue

        if attendance_date < from_date or attendance_date > to_date:
            continue

        # ---------------------------------------------------------------
        # Read values
        # ---------------------------------------------------------------

        in_value = ws.cell(row_number, 2).value
        out_value = ws.cell(row_number, 3).value

        exception_value = ws.cell(row_number, 9).value
        exception_text = (
            str(exception_value).strip()
            if exception_value not in (None, "")
            else "-"
        )

        regular_hours = _to_hours(
            ws.cell(row_number, 10).value
        )

        overtime_hours = _to_hours(
            ws.cell(row_number, 11).value
        )

        total_hours = _to_hours(
            ws.cell(row_number, 12).value
        )

        # Normalize Excel datetime values to time.
        if isinstance(in_value, datetime.datetime):
            in_value = in_value.time()

        if isinstance(out_value, datetime.datetime):
            out_value = out_value.time()

        rows_out.append(
            {
                "employee_name_raw": current_name,
                "device_id": current_device_id,
                "department_raw": current_department,
                "attendance_date": attendance_date,
                "in_time": (
                    in_value
                    if isinstance(in_value, datetime.time)
                    else None
                ),
                "out_time": (
                    out_value
                    if isinstance(out_value, datetime.time)
                    else None
                ),
                "exception": exception_text,
                "regular_hours": regular_hours,
                "overtime_hours": overtime_hours,
                "total_hours": total_hours,
            }
        )

    wb.close()

    return rows_out


# ---------------------------------------------------------------------------
# Controller
# ---------------------------------------------------------------------------

class BiometricAttendanceImport(Document):

    @frappe.whitelist()
    def parse_file(self):
        """
        Step 1:
        Read the attached Excel file, match Employees, and populate
        the preview table.

        No Attendance records are created here.
        """

        try:
            self._parse_file()

        except Exception:
            traceback = frappe.get_traceback()

            frappe.log_error(
                traceback,
                "Biometric Attendance Import - parse_file failed",
            )

            frappe.throw(
                (
                    "<b>Parse failed. Full traceback below:</b>"
                    "<pre style='white-space:pre-wrap;"
                    "font-size:11px'>{0}</pre>"
                ).format(
                    escape_html(traceback)
                ),
                title="Parse Error",
            )

    def _parse_file(self):

        # ---------------------------------------------------------------
        # Validate file
        # ---------------------------------------------------------------

        if not self.biometric_file:
            frappe.throw(
                "Attach the biometric export file first."
            )

        if not self.attendance_from_date:
            frappe.throw(
                "Attendance From Date is required."
            )

        if not self.attendance_to_date:
            frappe.throw(
                "Attendance To Date is required."
            )

        # ---------------------------------------------------------------
        # Find File document
        # ---------------------------------------------------------------

        file_doc = frappe.get_doc(
            "File",
            {
                "file_url": self.biometric_file,
            },
        )

        file_path = file_doc.get_full_path()

        if not file_path:
            frappe.throw(
                "Could not determine the path of the attached file."
            )

        # ---------------------------------------------------------------
        # Parse workbook
        # ---------------------------------------------------------------

        rows = _parse_workbook(
            file_path,
            self.attendance_from_date,
            self.attendance_to_date,
        )

        # ---------------------------------------------------------------
        # Build Employee device ID map
        #
        # Example:
        # {
        #     "1001": "HR-EMP-00001",
        #     "1002": "HR-EMP-00002"
        # }
        # ---------------------------------------------------------------

        employee_rows = frappe.get_all(
            "Employee",
            filters={
                "attendance_device_id": ["is", "set"],
            },
            fields=[
                "name",
                "attendance_device_id",
            ],
        )

        device_map = {}

        for employee_row in employee_rows:
            device_id = str(
                employee_row.attendance_device_id
            ).strip()

            if device_id:
                device_map[device_id] = employee_row.name

        # ---------------------------------------------------------------
        # Clear previous preview rows
        # ---------------------------------------------------------------

        self.set("import_details", [])

        # ---------------------------------------------------------------
        # Populate preview
        # ---------------------------------------------------------------

        for row in rows:

            device_id = str(
                row["device_id"]
            ).strip()

            employee = device_map.get(device_id)

            status = determine_status(
                row["exception"],
                row["regular_hours"],
                row["overtime_hours"],
                row["total_hours"],
            )

            self.append(
                "import_details",
                {
                    "employee_name_raw": row["employee_name_raw"],
                    "device_id": device_id,
                    "employee": employee,
                    "department_raw": row["department_raw"],
                    "attendance_date": row["attendance_date"],
                    "in_time": row["in_time"],
                    "out_time": row["out_time"],
                    "exception": row["exception"],
                    "regular_hours": row["regular_hours"],
                    "overtime_hours": row["overtime_hours"],
                    "total_hours": row["total_hours"],
                    "determined_status": status,
                    "result": (
                        "Pending"
                        if employee
                        else "Skipped - No Employee Match"
                    ),
                    "remarks": (
                        ""
                        if employee
                        else (
                            "No Employee found with this "
                            "Attendance Device ID"
                        )
                    ),
                },
            )

        # ---------------------------------------------------------------
        # Summary
        # ---------------------------------------------------------------

        self.total_rows_parsed = len(rows)

        unmatched = sum(
            1
            for detail in self.import_details
            if not detail.employee
        )

        self.status = "Parsed"

        self.import_log = (
            "Parsed {0} rows.\n"
            "{1} row(s) had no matching Employee "
            "(check Attendance Device ID on the Employee master).\n\n"
            "Review the table below, then click "
            "'Create Draft Attendance'."
        ).format(
            len(rows),
            unmatched,
        )

        self.save()

        frappe.msgprint(
            "Parsed {0} rows. Review the table, "
            "then create Attendance.".format(len(rows))
        )

    # -------------------------------------------------------------------
    # Create Attendance
    # -------------------------------------------------------------------

    @frappe.whitelist()
    def create_attendance(self):
        """
        Step 2:
        Walk through the reviewed preview table and create Attendance
        records as drafts.

        Existing Attendance records are skipped.
        Unmatched employees are skipped.
        """

        try:
            self._create_attendance()

        except Exception:
            traceback = frappe.get_traceback()

            frappe.log_error(
                traceback,
                "Biometric Attendance Import - create_attendance failed",
            )

            frappe.throw(
                (
                    "<b>Create Attendance failed. "
                    "Full traceback below:</b>"
                    "<pre style='white-space:pre-wrap;"
                    "font-size:11px'>{0}</pre>"
                ).format(
                    escape_html(traceback)
                ),
                title="Create Attendance Error",
            )

    def _create_attendance(self):

        created = 0
        skipped = 0
        errors = 0

        log_lines = []

        for row in self.import_details:

            # -----------------------------------------------------------
            # Already created during a previous run
            # -----------------------------------------------------------

            if row.result == "Created":
                continue

            # -----------------------------------------------------------
            # No Employee
            # -----------------------------------------------------------

            if not row.employee:

                row.result = "Skipped - No Employee Match"

                if not row.remarks:
                    row.remarks = (
                        "No Employee found with this "
                        "Attendance Device ID"
                    )

                skipped += 1

                continue

            # -----------------------------------------------------------
            # Validate date
            # -----------------------------------------------------------

            if not row.attendance_date:

                row.result = "Error"
                row.remarks = "Attendance Date is missing."

                errors += 1

                log_lines.append(
                    "Row {0} ({1}): Attendance Date is missing.".format(
                        row.idx,
                        row.employee,
                    )
                )

                continue

            # -----------------------------------------------------------
            # Prevent duplicate Attendance
            # -----------------------------------------------------------

            existing_attendance = frappe.db.exists(
                "Attendance",
                {
                    "employee": row.employee,
                    "attendance_date": row.attendance_date,
                    "docstatus": ["!=", 2],
                },
            )

            if existing_attendance:

                row.result = "Skipped - Duplicate"

                row.remarks = (
                    "Attendance already exists: {0}"
                ).format(
                    existing_attendance
                )

                skipped += 1

                continue

            # -----------------------------------------------------------
            # Create Attendance
            # -----------------------------------------------------------

            try:

                attendance = frappe.get_doc(
                    {
                        "doctype": "Attendance",
                        "employee": row.employee,
                        "attendance_date": row.attendance_date,
                        "status": row.determined_status,
                        "company": self.company,
                    }
                )

                # This creates a Draft Attendance because we do not submit it.
                attendance.insert(
                    ignore_permissions=True
                )

                row.result = "Created"
                row.remarks = attendance.name

                created += 1

            except Exception as exc:

                row.result = "Error"

                row.remarks = str(exc)[:140]

                errors += 1

                log_lines.append(
                    "Row {0} ({1}): {2}".format(
                        row.idx,
                        row.employee,
                        exc,
                    )
                )

            # -----------------------------------------------------------
            # Periodic commit for large imports
            # -----------------------------------------------------------

            processed = created + skipped + errors

            if processed > 0 and processed % 200 == 0:
                frappe.db.commit()

        # ---------------------------------------------------------------
        # Update import summary
        # ---------------------------------------------------------------

        self.total_created = created
        self.total_skipped = skipped
        self.total_errors = errors

        self.status = (
            "Completed"
            if errors == 0
            else "Failed"
        )

        previous_log = self.import_log or ""

        create_log = (
            "\n\n"
            "Create run: {0} created, "
            "{1} skipped, "
            "{2} errors.\n"
        ).format(
            created,
            skipped,
            errors,
        )

        if log_lines:
            create_log += "\n".join(log_lines)

        self.import_log = previous_log + create_log

        self.save()

        frappe.msgprint(
            (
                "Created {0}, skipped {1}, errors {2}. "
                "All new records are drafts. "
                "Review and submit them from the Attendance list."
            ).format(
                created,
                skipped,
                errors,
            )
        )